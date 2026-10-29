"""Generates the parser's page-count corpus (tests/corpus.rs) with the libraries
real exporters use: openpyxl (XLSX), xlwt (XLS), odfpy (ODS), reportlab (PDF).

    uv run --no-project --with openpyxl --with xlwt --with odfpy --with reportlab \
        python -I tests/corpus/generate.py tests/corpus

Writes each file and manifest.json with the pages the definition gives it.
"""
import csv, io, json, math, sys
from pathlib import Path

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
manifest = []

def rows(n, cols=("sku", "description", "qty", "unit_price", "total")):
    yield list(cols)
    for i in range(1, n + 1):
        yield [f"SKU-{i:05d}", f"Item number {i}", str(i % 7 + 1), f"{(i % 50) + 0.99:.2f}", f"{((i % 7) + 1) * ((i % 50) + 0.99):.2f}"]

def add(name, fmt, pages, note):
    manifest.append({"file": name, "format": fmt, "pages": pages, "note": note})

# CSV: a shop export with a BOM and CRLF, 2,400 data rows (+ header = 2,401 rows, 3 pages).
buf = io.StringIO()
w = csv.writer(buf, lineterminator="\r\n")
for r in rows(2400): w.writerow(r)
(out / "shop-export.csv").write_bytes(b"\xef\xbb\xbf" + buf.getvalue().encode())
add("shop-export.csv", "csv", 3, "2,401 rows with header, BOM, CRLF")

# CSV: exactly 1,000 rows including header (1 page), semicolons and quoted newlines.
buf = io.StringIO()
w = csv.writer(buf, delimiter=";")
w.writerow(["id", "note"])
for i in range(999): w.writerow([i, "line one\nline two" if i % 100 == 0 else f"note {i}"])
(out / "semicolon.csv").write_text(buf.getvalue())
add("semicolon.csv", "csv", 1, "1,000 rows including header, ; delimiter, quoted newlines")

# TSV: 1,001 rows (2 pages).
(out / "report.tsv").write_text("\n".join("\t".join(r) for r in rows(1000)) + "\n")
add("report.tsv", "tsv", 2, "1,001 rows")

# XLSX with openpyxl: two visible sheets and a hidden one.
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Orders"
for r in rows(1500): ws.append(r)
ws2 = wb.create_sheet("Summary")
for r in rows(10): ws2.append(r)
hidden = wb.create_sheet("Lookup")
for r in rows(2000): hidden.append(r)
hidden.sheet_state = "hidden"
wb.save(out / "orders.xlsx")
add("orders.xlsx", "xlsx", 3, "Orders 1,501 rows (2) + Summary 11 rows (1); hidden 2,001-row sheet skipped")

# XLS with xlwt (BIFF8, as legacy systems export).
import xlwt
book = xlwt.Workbook()
sheet = book.add_sheet("Stock")
for y, r in enumerate(rows(2100)):
    for x, v in enumerate(r): sheet.write(y, x, v)
book.save(str(out / "stock.xls"))
add("stock.xls", "xls", 3, "2,101 rows")

# ODS with odfpy.
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table, TableRow, TableCell
from odf.text import P
doc = OpenDocumentSpreadsheet()
t = Table(name="Bookings")
for r in rows(999):
    tr = TableRow()
    for v in r:
        c = TableCell(valuetype="string"); c.addElement(P(text=v)); tr.addElement(c)
    t.addElement(tr)
doc.spreadsheet.addElement(t)
doc.save(str(out / "bookings.ods"))
add("bookings.ods", "ods", 1, "1,000 rows including header")

# PDF invoices with reportlab: 1 page and 7 pages.
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
for name, n in (("invoice-1-page.pdf", 1), ("statement-7-pages.pdf", 7)):
    c = canvas.Canvas(str(out / name), pagesize=A4)
    for p in range(n):
        c.setFont("Helvetica", 11)
        c.drawString(72, 800, f"Acme Ltd. Statement page {p + 1} of {n}")
        for i in range(40): c.drawString(72, 770 - i * 16, f"SKU-{p * 40 + i:05d}  Item  2 x 9.99  19.98")
        c.showPage()
    c.save()
    add(name, "pdf", n, f"{n} pages, text layer, reportlab")

# Text attachments.
(out / "calendar.ics").write_text("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nSUMMARY:Booking 1001\r\nDTSTART:20261101T100000Z\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n")
add("calendar.ics", "text", 1, "calendar invite")
(out / "payload.json").write_text(json.dumps({"order": 1001, "items": [{"sku": "A-1", "qty": 2}]}, indent=2) + "\n")
add("payload.json", "text", 1, "JSON export")

(out / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
print(len(manifest), "files")
